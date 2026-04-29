"""
Zoho Mail Extractor
Fetches up to 1,000 emails and saves them to zoho_emails.xlsx
Run: python extractor.py
"""
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import config

def get_access_token():
    if getattr(config, "ACCESS_TOKEN", ""):
        return config.ACCESS_TOKEN
    resp = requests.post(f"{config.ZOHO_ACCOUNTS_URL}/oauth/v2/token", data={
        "refresh_token": config.REFRESH_TOKEN,
        "client_id": config.CLIENT_ID,
        "client_secret": config.CLIENT_SECRET,
        "grant_type": "refresh_token",
    })
    data = resp.json()
    if "access_token" not in data:
        raise Exception(f"Failed to get access token: {data}")
    return data["access_token"]

def get_account_id(token):
    resp = requests.get(f"{config.ZOHO_MAIL_API}/accounts", headers={"Authorization": f"Zoho-oauthtoken {token}"})
    data = resp.json()
    accounts = data.get("data", [])
    if not accounts:
        raise Exception("No Zoho Mail accounts found.")
    return accounts[0]["accountId"]

def fetch_emails(token, account_id):
    emails = []
    start = 1
    batch = 200
    limit_total = config.EMAIL_LIMIT

    label = f"up to {limit_total}" if limit_total else "ALL"
    print(f"Fetching {label} emails...")

    while True:
        fetch_count = batch if not limit_total else min(batch, limit_total - len(emails))
        params = {
            "start": start,
            "limit": fetch_count,
            "sortorder": "false",  # newest first
        }
        url = f"{config.ZOHO_MAIL_API}/accounts/{account_id}/messages/view"
        resp = requests.get(url, headers={"Authorization": f"Zoho-oauthtoken {token}"}, params=params)
        data = resp.json()

        batch_emails = data.get("data", [])
        if not batch_emails:
            print(f"  No more emails found. Total: {len(emails)}")
            break

        emails.extend(batch_emails)
        print(f"  Fetched {len(emails)} emails so far...")

        if len(batch_emails) < fetch_count:
            break

        if limit_total and len(emails) >= limit_total:
            break

        start += fetch_count

    return emails

def get_email_body(token, account_id, folder_id, message_id):
    url = f"{config.ZOHO_MAIL_API}/accounts/{account_id}/folders/{folder_id}/messages/{message_id}/content"
    resp = requests.get(url, headers={"Authorization": f"Zoho-oauthtoken {token}"})
    data = resp.json()
    content = data.get("data", {})
    body = content.get("content", "")
    import re
    body = re.sub(r"<[^>]+>", " ", body)
    body = re.sub(r"\s+", " ", body).strip()
    return body[:5000]

def save_to_excel(emails, token, account_id):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Emails"

    headers = ["#", "Date", "From", "To", "CC", "Subject", "Body", "Attachments"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 25
    col_widths = [5, 20, 30, 30, 25, 40, 60, 12]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width

    print("\nProcessing emails and writing to Excel...")

    for idx, email in enumerate(emails, 1):
        print(f"  Writing email {idx}/{len(emails)}...", end="\r")

        # Parse date
        sent_time = email.get("sentDateInGMT", "") or email.get("receivedTime", "")
        try:
            ts = int(sent_time) / 1000
            date_str = datetime.utcfromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S UTC")
        except Exception:
            date_str = sent_time

        from_addr = email.get("fromAddress", "")
        to_addr = email.get("toAddress", "")
        cc_addr = email.get("ccAddress", "")
        subject = email.get("subject", "(No Subject)")
        has_attach = email.get("hasAttachment", "0")
        attach_count = email.get("attachmentCount", 0) if has_attach == "1" else 0
        message_id = email.get("messageId", "")
        folder_id = email.get("folderId", "")

        # Fetch body
        body = ""
        if message_id and folder_id:
            try:
                body = get_email_body(token, account_id, folder_id, message_id)
            except Exception:
                body = "[Could not fetch body]"

        row = [idx, date_str, from_addr, to_addr, cc_addr, subject, body, attach_count]
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=idx + 1, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Alternate row shading
        if idx % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=idx + 1, column=col).fill = PatternFill(
                    start_color="EBF3FB", end_color="EBF3FB", fill_type="solid"
                )

    output = "zoho_emails.xlsx"
    wb.save(output)
    print(f"\n\n✅ Done! Saved {len(emails)} emails to '{output}'")
    return output

def main():
    if not config.REFRESH_TOKEN and not getattr(config, "ACCESS_TOKEN", ""):
        print("❌ No token found in config.py")
        print("   Run 'python3 setup.py' first to get your Refresh Token.")
        return

    print("=== Zoho Mail Extractor ===\n")
    print("Authenticating...")
    token = get_access_token()
    print("✅ Authenticated\n")

    print("Fetching account info...")
    account_id = get_account_id(token)
    print(f"✅ Account: abshar@deltainstitutions.com (ID: {account_id})\n")

    emails = fetch_emails(token, account_id)
    save_to_excel(emails, token, account_id)

if __name__ == "__main__":
    main()
